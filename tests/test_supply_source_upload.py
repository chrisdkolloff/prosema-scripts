"""Supply-source upload parse, template generate, and resolve branch."""

from __future__ import annotations

from decimal import Decimal
from unittest.mock import patch

import pytest
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db import engine
from app.models import (
    Supplier,
    SupplySourceRow,
    SupplySourceRun,
    SupplySourceUpload,
    WeclappArticle,
    WeclappSupplySource,
    WeclappSupplySourceLink,
    WeclappSupplySourcePrice,
    WeclappUnit,
)
from app.supply_source_resolve import run_resolve
from app.supply_source_runs import create_upload_run
from app.supply_source_templates import (
    DEFAULT_COLUMNS,
    generate_template_xlsx_for_user,
    get_or_create_active_template,
)
from app.supply_source_upload import (
    SupplySourceParseError,
    parse_listenpreis,
    parse_upload_bytes,
)

PLAIN_USER = {
    "oid": "user-oid-ss-upload",
    "name": "Dennis",
    "email": "user@example.com",
    "roles": ["user"],
}


@pytest.fixture
def db_session():
    connection = engine.connect()
    trans = connection.begin()
    session = Session(
        bind=connection,
        join_transaction_mode="create_savepoint",
        expire_on_commit=False,
    )
    try:
        yield session
    finally:
        session.close()
        trans.rollback()
        connection.close()


def _now():
    from datetime import UTC, datetime

    return datetime.now(UTC)


def _supplier(db: Session, number: str = "19997") -> Supplier:
    existing = db.scalars(select(Supplier).where(Supplier.supplier_number == number)).first()
    if existing:
        return existing
    row = Supplier(
        supplier_number=number,
        weclapp_party_id=f"party-upload-{number}",
        name="Upload-Testlieferant",
        einkaufswaehrung="EUR",
        default_kurs=Decimal("0.93"),
        default_aufschlag=Decimal("0.50"),
        default_verkaufswaehrung="CHF",
        is_active=True,
    )
    db.add(row)
    db.flush()
    return row


def _unit(db: Session) -> None:
    if db.get(WeclappUnit, "3566") is None:
        db.add(
            WeclappUnit(
                weclapp_id="3566",
                name="Stk.",
                description="Stück",
                last_seen_at=_now(),
            )
        )
    if db.get(WeclappUnit, "4302") is None:
        db.add(
            WeclappUnit(
                weclapp_id="4302",
                name="Btl.",
                description="Beutel",
                last_seen_at=_now(),
            )
        )
    db.flush()


def _csv(rows: list[list[str]]) -> bytes:
    lines = [";".join(r) for r in rows]
    return ("\n".join(lines) + "\n").encode("utf-8")


def _header() -> list[str]:
    return [c["label"] for c in DEFAULT_COLUMNS]


def test_parse_listenpreis_comma_and_dot():
    assert parse_listenpreis("1.234,56") == Decimal("1234.56")
    assert parse_listenpreis("1'234.56") == Decimal("1234.56")
    assert parse_listenpreis("1234,56") == Decimal("1234.56")
    assert parse_listenpreis("1234.56") == Decimal("1234.56")


def test_non_numeric_listenpreis_kept_as_priceless(db_session):
    _unit(db_session)
    data = _csv(
        [
            _header(),
            ["SAN-ASK", "Tonne", "auf Anfrage", "", "t", "", "", ""],
            ["SAN-OK", "Teil", "10,00", "", "Stk.", "", "", ""],
        ]
    )
    parsed = parse_upload_bytes(
        db_session, data, filename="x.csv", columns=DEFAULT_COLUMNS
    )
    by_san = {r.supplier_article_number: r for r in parsed.rows}
    assert len(parsed.rows) == 2
    assert by_san["SAN-ASK"].listenpreis is None
    assert by_san["SAN-ASK"].listenpreis_raw == "auf Anfrage"
    assert parsed.row_errors == []
    assert any("auf Anfrage" in m for m in parsed.priceless)
    assert by_san["SAN-OK"].listenpreis == Decimal("10.00")


def test_duplicate_san_rejects_before_rows(db_session):
    supplier = _supplier(db_session)
    _unit(db_session)
    data = _csv(
        [
            _header(),
            ["DUP", "A", "1,00", "", "Stk.", "", "", ""],
            ["DUP", "B", "2,00", "", "Stk.", "", "", ""],
        ]
    )
    with pytest.raises(SupplySourceParseError, match="Doppelte Lieferantenartikelnummer"):
        parse_upload_bytes(
            db_session, data, filename="x.csv", columns=DEFAULT_COLUMNS
        )
    before_rows = db_session.scalar(select(func.count()).select_from(SupplySourceRow))
    before_uploads = db_session.scalar(select(func.count()).select_from(SupplySourceUpload))
    with pytest.raises(Exception, match="Doppelte"):
        create_upload_run(
            db_session,
            supplier_id=supplier.id,
            filename="x.csv",
            content=data,
            user=PLAIN_USER,
        )
    assert db_session.scalar(select(func.count()).select_from(SupplySourceRow)) == before_rows
    assert db_session.scalar(select(func.count()).select_from(SupplySourceUpload)) == before_uploads


def test_unknown_unit_reports_and_null(db_session):
    _unit(db_session)
    data = _csv(
        [
            _header(),
            ["SAN-1", "Teil", "10,00", "", "XXEINHEIT-UNBEKANNT", "", "", ""],
        ]
    )
    parsed = parse_upload_bytes(
        db_session, data, filename="x.csv", columns=DEFAULT_COLUMNS
    )
    assert parsed.rows[0].unit_id is None
    assert parsed.rows[0].file_unit_id is None
    assert parsed.rows[0].unit_raw == "XXEINHEIT-UNBEKANNT"
    assert parsed.unmatched_units[0]["value"] == "XXEINHEIT-UNBEKANNT"


def test_parse_matches_description_casefold(db_session):
    _unit(db_session)
    data = _csv(
        [
            _header(),
            ["SAN-STK", "Teil", "10,00", "", "stück", "", "", ""],
            ["SAN-BTL", "Beutelware", "10,00", "", " Beutel ", "", "", ""],
        ]
    )
    parsed = parse_upload_bytes(
        db_session, data, filename="x.csv", columns=DEFAULT_COLUMNS
    )
    by_san = {r.supplier_article_number: r for r in parsed.rows}
    assert by_san["SAN-STK"].unit_id == "3566"
    assert by_san["SAN-STK"].unit_raw == "stück"
    assert by_san["SAN-BTL"].unit_id == "4302"
    assert by_san["SAN-BTL"].unit_raw == "Beutel"
    assert parsed.unmatched_units == []


def test_parse_uses_supplier_unit_alias(db_session):
    from app.models import SupplierUnitAlias
    from app.weclapp_units import fold_unit_word

    supplier = _supplier(db_session, "19990")
    _unit(db_session)
    db_session.add(
        SupplierUnitAlias(
            supplier_id=supplier.id,
            supplier_word="xyztonne",
            word_key=fold_unit_word("xyztonne"),
            unit_id="3566",
        )
    )
    db_session.flush()
    data = _csv(
        [
            _header(),
            ["SAN-T", "Tonne", "10,00", "", "XYZtonne", "", "", ""],
        ]
    )
    parsed = parse_upload_bytes(
        db_session,
        data,
        filename="x.csv",
        columns=DEFAULT_COLUMNS,
        supplier_id=supplier.id,
    )
    assert parsed.rows[0].unit_id == "3566"
    assert parsed.rows[0].unit_raw == "XYZtonne"
    assert parsed.unmatched_units == []


def test_upload_zero_ss_create_or_unmatched(db_session):
    supplier = _supplier(db_session, "19996")
    _unit(db_session)
    db_session.add(
        WeclappArticle(
            weclapp_article_id="upl-a-ean",
            article_number="999.010.5010",
            name="Artikel",
            ean="5901234123457",
            unit_id="3566",
            weclapp_version="1",
            last_seen_at=_now(),
        )
    )
    db_session.flush()
    data = _csv(
        [
            _header(),
            ["NEW-MATCH", "Neu", "5,00", "5901234123457", "Stk.", "", "", ""],
            ["NEW-MISS", "Ohne", "5,00", "", "Stk.", "", "", ""],
        ]
    )
    with patch("app.jobs.enqueue") as enqueue:
        from uuid import uuid4

        from app.models import Job

        job = Job(
            id=uuid4(),
            job_type="supply_source_resolve",
            payload={},
            status="queued",
            created_by_oid="x",
            created_by_name="x",
        )
        db_session.add(job)
        db_session.flush()
        enqueue.return_value = job
        run = create_upload_run(
            db_session,
            supplier_id=supplier.id,
            filename="neu.csv",
            content=data,
            user=PLAIN_USER,
        )
    run_resolve(db_session, run, oid="x", skip_index=True)
    rows = {
        r.supplier_article_number: r
        for r in db_session.scalars(
            select(SupplySourceRow).where(SupplySourceRow.run_id == run.id)
        )
    }
    assert rows["NEW-MATCH"].row_intent == "create"
    assert rows["NEW-MATCH"].match_tier == 3
    assert rows["NEW-MISS"].match_status == "unmatched"


def test_upload_name_divergence_defaults_weclapp(db_session):
    supplier = _supplier(db_session, "19995")
    _unit(db_session)
    now = _now()
    db_session.add(
        WeclappSupplySource(
            weclapp_id="upl-ss-1",
            supplier_party_id=supplier.weclapp_party_id,
            supplier_number=supplier.supplier_number,
            supplier_article_number="L-1",
            name="Altname",
            weclapp_version="1",
            last_seen_at=now,
            unit_id="3566",
        )
    )
    db_session.add(
        WeclappSupplySource(
            weclapp_id="upl-ss-2",
            supplier_party_id=supplier.weclapp_party_id,
            supplier_number=supplier.supplier_number,
            supplier_article_number="L-2",
            name="Gleich",
            weclapp_version="1",
            last_seen_at=now,
            unit_id="3566",
        )
    )
    db_session.add(
        WeclappSupplySourcePrice(
            supply_source_weclapp_id="upl-ss-1",
            price=Decimal("10.0000"),
            currency_code="EUR",
        )
    )
    db_session.add(
        WeclappSupplySourcePrice(
            supply_source_weclapp_id="upl-ss-2",
            price=Decimal("10.0000"),
            currency_code="EUR",
        )
    )
    db_session.add(
        WeclappArticle(
            weclapp_article_id="upl-a1",
            article_number="999.010.6010",
            name="A",
            unit_id="3566",
            weclapp_version="1",
            last_seen_at=now,
        )
    )
    db_session.add(
        WeclappArticle(
            weclapp_article_id="upl-a2",
            article_number="999.010.6020",
            name="B",
            unit_id="3566",
            weclapp_version="1",
            last_seen_at=now,
        )
    )
    db_session.add(
        WeclappSupplySourceLink(
            supply_source_weclapp_id="upl-ss-1",
            weclapp_article_id="upl-a1",
            article_number="999.010.6010",
            supplier_party_id=supplier.weclapp_party_id,
        )
    )
    db_session.add(
        WeclappSupplySourceLink(
            supply_source_weclapp_id="upl-ss-2",
            weclapp_article_id="upl-a2",
            article_number="999.010.6020",
            supplier_party_id=supplier.weclapp_party_id,
        )
    )
    db_session.flush()
    data = _csv(
        [
            _header(),
            ["L-1", "Neuname", "10,00", "", "Stk.", "", "", ""],
            ["L-2", "Gleich", "12,00", "", "Stk.", "", "", ""],
        ]
    )
    with patch("app.jobs.enqueue") as enqueue:
        from uuid import uuid4

        from app.models import Job

        job = Job(
            id=uuid4(),
            job_type="supply_source_resolve",
            payload={},
            status="queued",
            created_by_oid="x",
            created_by_name="x",
        )
        db_session.add(job)
        db_session.flush()
        enqueue.return_value = job
        run = create_upload_run(
            db_session,
            supplier_id=supplier.id,
            filename="div.csv",
            content=data,
            user=PLAIN_USER,
        )
    run_resolve(db_session, run, oid="x", skip_index=True)
    rows = {
        r.supplier_article_number: r
        for r in db_session.scalars(
            select(SupplySourceRow).where(SupplySourceRow.run_id == run.id)
        )
    }
    assert rows["L-1"].name == "Altname"
    assert rows["L-1"].template_name == "Neuname"
    assert rows["L-1"].field_overrides.get("name") == "weclapp"
    assert rows["L-1"].row_intent == "price_only"
    assert rows["L-2"].listenpreis == Decimal("12.00")
    assert rows["L-2"].row_intent == "price_only"
    assert not rows["L-2"].field_overrides


def test_generate_then_upload_roundtrip(db_session):
    supplier = _supplier(db_session, "19994")
    _unit(db_session)
    now = _now()
    db_session.add(
        WeclappSupplySource(
            weclapp_id="upl-ss-rt",
            supplier_party_id=supplier.weclapp_party_id,
            supplier_number=supplier.supplier_number,
            supplier_article_number="RT-1",
            name="Spiegel",
            ean="4000000000001",
            weclapp_version="1",
            last_seen_at=now,
            unit_id="3566",
        )
    )
    db_session.add(
        WeclappSupplySourcePrice(
            supply_source_weclapp_id="upl-ss-rt",
            price=Decimal("8.5000"),
            currency_code="EUR",
        )
    )
    db_session.add(
        WeclappArticle(
            weclapp_article_id="upl-rt-a",
            article_number="999.010.7010",
            name="A",
            unit_id="3566",
            weclapp_version="1",
            last_seen_at=now,
        )
    )
    db_session.add(
        WeclappSupplySourceLink(
            supply_source_weclapp_id="upl-ss-rt",
            weclapp_article_id="upl-rt-a",
            article_number="999.010.7010",
            supplier_party_id=supplier.weclapp_party_id,
        )
    )
    db_session.flush()
    _tpl, xlsx = generate_template_xlsx_for_user(db_session, supplier, user=PLAIN_USER)
    with patch("app.jobs.enqueue") as enqueue:
        from uuid import uuid4

        from app.models import Job

        job = Job(
            id=uuid4(),
            job_type="supply_source_resolve",
            payload={},
            status="queued",
            created_by_oid="x",
            created_by_name="x",
        )
        db_session.add(job)
        db_session.flush()
        enqueue.return_value = job
        run = create_upload_run(
            db_session,
            supplier_id=supplier.id,
            filename="vorlage.xlsx",
            content=xlsx,
            user=PLAIN_USER,
        )
    run_resolve(db_session, run, oid="x", skip_index=True)
    row = db_session.scalars(
        select(SupplySourceRow).where(SupplySourceRow.run_id == run.id)
    ).one()
    assert row.match_tier == 1
    assert row.row_intent == "price_only"
    assert not row.field_overrides


def test_blank_template_has_headers_only(db_session):
    from io import BytesIO

    from openpyxl import load_workbook

    _tpl, xlsx = generate_template_xlsx_for_user(db_session, user=PLAIN_USER)
    wb = load_workbook(BytesIO(xlsx))
    ws = wb.active
    assert ws["A1"].value == "Lieferantenartikelnummer"
    assert ws["C1"].value == "Listenpreis"
    assert ws.max_row == 2
    assert all((ws.cell(row=2, column=i).value in (None, "")) for i in range(1, 8))


def test_get_or_create_template_versions(db_session):
    t1 = get_or_create_active_template(db_session, user=PLAIN_USER)
    t2 = get_or_create_active_template(db_session, user=PLAIN_USER)
    assert t1.id == t2.id
    assert t1.version == 1
